#!/usr/bin/env python3
"""Extract numbered and obligation-bearing requirements from an RFP into CSV.

Supports .txt, .md, .docx (python-docx if installed), and .pdf (pypdf if
installed). When an optional library is missing, the script says so and exits
with code 3 rather than failing silently. No network access is used.

Detection:
  * Numbered clauses:  1.2.3 / 4.3 / R-014 / REQ 12 / (a) / (ii) / A.1
  * Obligation words:  shall, must, is required to, required, will provide,
                       mandatory, should, expected to, preferred  (M/D classification)
  * Tables in .docx are read row by row; the first cell is used as the client ID
    when it looks like an identifier.

Output columns (match assets/compliance-matrix-columns.md, cols 1-9 subset):
  req_id, client_id, section, page, text, mdo_hint, keyword, source_file

Usage:
    python3 -I extract_requirements.py rfp.pdf --out requirements.csv
    python3 -I extract_requirements.py rfp.docx appendix-b.docx --out reqs.csv --min-words 4
    python3 -I extract_requirements.py rfp.md --dry-run

Exit codes: 0 ok; 2 file not found/unsupported; 3 missing optional library.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path

CLAUSE_RE = re.compile(
    r"^\s*(?:"
    r"(?P<num>\d+(?:\.\d+){0,4})[.)]?\s+"          # 1.2.3 or 4.3
    r"|(?P<req>(?:REQ|R|FR|NFR|TR|SR)[-\s]?\d+[A-Za-z]?)[:.)]?\s+"  # R-014, REQ 12
    r"|\((?P<alpha>[a-z]{1,2}|[ivx]{1,5})\)\s+"      # (a) (ii)
    r"|(?P<letnum>[A-Z]\.\d+(?:\.\d+)*)[.)]?\s+"     # A.1.2
    r")"
)
HEADING_RE = re.compile(r"^\s*(?:\d+(?:\.\d+)*[.)]?\s+)?([A-Z][A-Za-z0-9 ,&/()-]{3,80})\s*$")
MANDATORY_RE = re.compile(
    r"\b(shall|must|is required to|are required to|required|will provide|mandatory|minimum requirement)\b",
    re.IGNORECASE,
)
DESIRABLE_RE = re.compile(
    r"\b(should|expected to|preferred|desirable|where possible|advantageous)\b", re.IGNORECASE
)
OPTIONAL_RE = re.compile(r"\b(may|optional|nice to have|could)\b", re.IGNORECASE)
BOILERPLATE_RE = re.compile(
    r"^(table of contents|page \d+|confidential|copyright|all rights reserved)", re.IGNORECASE
)


@dataclass
class Req:
    client_id: str
    section: str
    page: str
    text: str
    mdo_hint: str
    keyword: str
    source_file: str


def read_txt(path: Path) -> list[tuple[str, str]]:
    """Return list of (page_label, text) blocks. Plain text has one block."""
    return [("1", path.read_text(encoding="utf-8", errors="replace"))]


def read_docx(path: Path) -> list[tuple[str, str]]:
    try:
        import docx  # type: ignore
    except ImportError:
        print("error: python-docx not installed; convert the .docx to .txt or install python-docx", file=sys.stderr)
        sys.exit(3)
    d = docx.Document(str(path))
    lines: list[str] = []
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            style = (para.style.name or "").lower() if para.style is not None else ""
            lines.append(("## " + t) if style.startswith("heading") else t)
    for ti, table in enumerate(d.tables, start=1):
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            cells = [c for c in cells if c]
            if not cells:
                continue
            if len(cells) >= 2 and re.match(r"^[A-Za-z]{0,4}[-\s]?\d+(\.\d+)*[A-Za-z]?$", cells[0]):
                lines.append(f"{cells[0]} {' '.join(cells[1:])}")
            else:
                lines.append(" ".join(cells))
    return [("docx", "\n".join(lines))]


def read_pdf(path: Path) -> list[tuple[str, str]]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        print("error: pypdf not installed; export the PDF to text or install pypdf", file=sys.stderr)
        sys.exit(3)
    reader = PdfReader(str(path))
    blocks: list[tuple[str, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception as e:  # noqa: BLE001 - report and continue
            txt = ""
            print(f"warning: page {i} text extraction failed: {e}", file=sys.stderr)
        blocks.append((str(i), txt))
    if not any(t.strip() for _, t in blocks):
        print("warning: no extractable text (scanned PDF?). OCR is out of scope for this script.", file=sys.stderr)
    return blocks


READERS = {".txt": read_txt, ".md": read_txt, ".docx": read_docx, ".pdf": read_pdf}


def mdo_hint(text: str) -> tuple[str, str]:
    m = MANDATORY_RE.search(text)
    if m:
        return "M", m.group(1).lower()
    m = DESIRABLE_RE.search(text)
    if m:
        return "D", m.group(1).lower()
    m = OPTIONAL_RE.search(text)
    if m:
        return "O", m.group(1).lower()
    return "", ""


def join_wrapped_lines(text: str) -> list[str]:
    """Merge PDF-wrapped lines into sentences: a line that does not end a sentence
    and whose successor does not start a clause is appended to the previous line."""
    out: list[str] = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            out.append("")
            continue
        if out and out[-1] and not re.search(r"[.;:]\s*$", out[-1]) and not CLAUSE_RE.match(line) \
                and not line.startswith("## ") and not HEADING_RE.match(line):
            out[-1] = out[-1] + " " + line.strip()
        else:
            out.append(line)
    return out


def extract(blocks: list[tuple[str, str]], source: str, min_words: int, numbered_only: bool) -> list[Req]:
    reqs: list[Req] = []
    section = ""
    for page, text in blocks:
        for line in join_wrapped_lines(text):
            s = line.strip()
            if not s or BOILERPLATE_RE.match(s):
                continue
            if s.startswith("## "):
                section = s[3:].strip()
                continue
            cm = CLAUSE_RE.match(s)
            client_id = ""
            body = s
            if cm:
                client_id = next(g for g in (cm.group("num"), cm.group("req"), cm.group("alpha"), cm.group("letnum")) if g)
                body = s[cm.end():].strip()
                # A short numbered line with no obligation word is a heading, not a requirement.
                if len(body.split()) <= 8 and not MANDATORY_RE.search(body) and not DESIRABLE_RE.search(body):
                    section = f"{client_id} {body}"
                    continue
            elif HEADING_RE.match(s) and len(s.split()) <= 10 and not MANDATORY_RE.search(s):
                section = s
                continue
            hint, kw = mdo_hint(body)
            if numbered_only and not client_id:
                continue
            if not client_id and not hint:
                continue
            if len(body.split()) < min_words:
                continue
            reqs.append(Req(client_id, section, page, body, hint, kw, source))
    return reqs


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("files", nargs="+", help="RFP files (.txt .md .docx .pdf)")
    p.add_argument("--out", default="requirements.csv", help="Output CSV path")
    p.add_argument("--min-words", type=int, default=5, help="Ignore candidate lines shorter than this")
    p.add_argument("--numbered-only", action="store_true", help="Keep only lines with a clause identifier")
    p.add_argument("--dry-run", action="store_true", help="Print summary; do not write CSV")
    args = p.parse_args(argv)

    all_reqs: list[Req] = []
    for f in args.files:
        path = Path(f)
        if not path.is_file():
            print(f"error: file not found: {path}", file=sys.stderr)
            return 2
        reader = READERS.get(path.suffix.lower())
        if reader is None:
            print(f"error: unsupported file type: {path.suffix} (use .txt .md .docx .pdf)", file=sys.stderr)
            return 2
        blocks = reader(path)
        all_reqs.extend(extract(blocks, path.name, args.min_words, args.numbered_only))

    # De-duplicate exact text repeats (headers/footers repeated per page)
    seen: set[str] = set()
    unique: list[Req] = []
    for r in all_reqs:
        key = r.text.lower()
        if key in seen:
            continue
        seen.add(key)
        unique.append(r)

    m = sum(1 for r in unique if r.mdo_hint == "M")
    d = sum(1 for r in unique if r.mdo_hint == "D")
    o = sum(1 for r in unique if r.mdo_hint == "O")
    print(f"candidates: {len(unique)}  (M={m} D={d} O={o} unclassified={len(unique)-m-d-o})", file=sys.stderr)

    if args.dry_run:
        for i, r in enumerate(unique[:25], start=1):
            print(f"REQ-{i:03d} | {r.client_id or '-'} | {r.mdo_hint or '-'} | {r.text[:90]}")
        if len(unique) > 25:
            print(f"... {len(unique) - 25} more")
        return 0

    with open(args.out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["req_id", "client_id", "section", "page", "text", "mdo_hint", "keyword", "source_file"])
        for i, r in enumerate(unique, start=1):
            w.writerow([f"REQ-{i:03d}", r.client_id, r.section, r.page, r.text, r.mdo_hint, r.keyword, r.source_file])
    print(f"wrote {args.out} ({len(unique)} rows). Review: merge split sentences, drop boilerplate, add table/appendix items.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
