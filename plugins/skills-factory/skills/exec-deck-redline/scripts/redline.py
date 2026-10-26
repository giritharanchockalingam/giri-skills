#!/usr/bin/env python3
"""Deck redline tooling: per-slide text ledger + unsourced-number list; apply verdicts back.

    python3 -I scripts/redline.py deck.pptx --ledger ledger.csv --unsourced unsourced-numbers.csv
    python3 -I scripts/redline.py deck.md   --ledger ledger.csv --unsourced unsourced-numbers.csv
    python3 -I scripts/redline.py --apply redline.md --ledger ledger.csv

Inputs: a .pptx (needs python-pptx: `pip install python-pptx`; see requirements-optional.txt) or a
.md/.txt export with `--- slide N ---` separators (first non-empty line of each block = title).
ledger.csv: slide, title, text, verdict, reason, numbers_found   (verdict/reason blank until applied)
unsourced-numbers.csv: slide, number, context — every number (>= 3 significant digits or carrying
$ / %) on a line that names no source (`source`, `per <doc>`, `[Sn]`, `(… 20xx)` count as sources).
--apply reads the redline table in redline.md (| Slide # | Original title | Verdict | Reason | …)
and fills verdict/reason for the matching slides; rows without a bracketed verdict are ignored.
--draft-only --redline-out redline.md (build mode): no model step — writes a mechanical redline on the
template's sections from the ledger alone: one row per slide, [TBD] + [SHARPEN] for slides with
unsourced figures (each figure becomes an owned to-be-baselined marker in the rewritten slide),
[SHARPEN] for a label title (<= 4 words, no figure, no verb of decision) or a body over 60 words,
[KEEP] otherwise; a missing decision ask in slides 1-2 is listed under "Slides to add". Labelled as a
draft; storyline untouched; the --apply step then fills the ledger from it.
Nothing is invented; the ledger is the deck's own text. Exit 1 on a missing file, an unsupported
extension, a text export without separators, or a pptx without python-pptx. Stdlib only, offline.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

NUMBER_RE = re.compile(r"(?<![\w.])([$€£]?)(-?\d[\d,]*(?:\.\d+)?)(%?)(?![\w])")
SOURCE_RE = re.compile(r"(?i:\bsources?\b)|\bper\s+[A-Z(]|\[S\d+\]|\(\s*[^()]*20\d\d\s*\)|(?i:\bsrc:)")
SEP_RE = re.compile(r"^---\s*slide\s+(\d+)\s*---\s*$", re.I)
DATE_RE = re.compile(r"\b\d{4}-\d{2}(-\d{2})?\b")
FOOTNOTE_RE = re.compile(r"^\s*(source|sources)\s*[:—-]", re.I)
VERDICT_RE = re.compile(r"\[(KEEP|SHARPEN|CUT|ADD|TBD)\]")
LEDGER_COLS = ["slide", "title", "text", "verdict", "reason", "numbers_found"]


def die(msg: str) -> None:
    print(f"error: {msg}", file=sys.stderr)
    sys.exit(1)


def sig_digits(num: str) -> int:
    return len(num.replace(",", "").replace("-", "").replace(".", "").lstrip("0"))


def numbers_in(text: str) -> list[str]:
    """Numbers worth tracing (>= 3 significant digits or $ / %); ISO dates are not numbers."""
    out = []
    for m in NUMBER_RE.finditer(DATE_RE.sub(" ", text)):
        sym, num, pct = m.groups()
        if sym or pct or sig_digits(num) >= 3:
            out.append(f"{sym}{num}{pct}")
    return out


def slides_from_text(text: str) -> list[tuple[int, str, str]]:
    slides, cur, buf = [], None, []
    for ln in text.splitlines():
        m = SEP_RE.match(ln.strip())
        if m:
            if cur is not None:
                slides.append((cur, buf))
            cur, buf = int(m.group(1)), []
        elif cur is not None:
            buf.append(ln.rstrip())
    if cur is not None:
        slides.append((cur, buf))
    if not slides:
        die("text export has no '--- slide N ---' separators")
    out = []
    for n, lines in slides:
        body = [ln for ln in lines if ln.strip()]
        title = body[0].lstrip("# ").strip() if body else ""
        out.append((n, title, "\n".join(body)))
    return out


def slides_from_pptx(path: Path) -> list[tuple[int, str, str]]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        die("reading .pptx needs python-pptx: pip install python-pptx  (or export the deck to "
            "deck.md with '--- slide N ---' separators)")
    prs = Presentation(str(path))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        lines, title = [], ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:  # noqa: BLE001
            title = ""
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if getattr(slide, "has_notes_slide", False) and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                lines.append(f"[notes] {notes}")
        if not title and lines:
            title = lines[0]
        out.append((n, title, "\n".join(lines)))
    if not out:
        die("pptx has no slides")
    return out


def build(deck: Path, ledger: Path, unsourced: Path, redline_out: Path | None = None) -> None:
    if not deck.is_file():
        die(f"file not found: {deck}")
    ext = deck.suffix.lower()
    if ext == ".pptx":
        slides = slides_from_pptx(deck)
    elif ext in (".md", ".txt"):
        slides = slides_from_text(deck.read_text(encoding="utf-8", errors="replace"))
    else:
        die(f"unsupported deck format {ext!r}: use .pptx, .md or .txt")
    rows, uns = [], []
    for n, title, text in slides:
        nums = []
        footnote = any(FOOTNOTE_RE.match(ln) for ln in text.splitlines())
        for ln in text.splitlines():
            found = numbers_in(ln)
            nums += found
            if found and not SOURCE_RE.search(ln) and not footnote:
                uns += [{"slide": n, "number": x, "context": ln.strip()[:160]} for x in found]
        rows.append({"slide": n, "title": title, "text": text, "verdict": "", "reason": "",
                     "numbers_found": "; ".join(nums)})
    ledger.parent.mkdir(parents=True, exist_ok=True)
    with ledger.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LEDGER_COLS); w.writeheader(); w.writerows(rows)
    with unsourced.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["slide", "number", "context"]); w.writeheader(); w.writerows(uns)
    if redline_out is not None:
        redline_out.parent.mkdir(parents=True, exist_ok=True)
        redline_out.write_text(draft_redline(rows, uns, deck.name), encoding="utf-8")
    print(f"ledger: {len(rows)} slides, {sum(len(r['numbers_found'].split('; ')) for r in rows if r['numbers_found'])} "
          f"numbers, {len(uns)} unsourced -> {ledger}, {unsourced}")


DECK_TBD = "[TO BE BASELINED — owner: Deck owner]"
DECISION_RE = re.compile(r"\b(decision|approve|approval|decide|ask)\b", re.I)


def draft_redline(rows: list[dict], uns: list[dict], deck_name: str) -> str:
    """Mechanical redline (no model step): checks from the rubric that the ledger alone can decide."""
    by_slide: dict[int, list[dict]] = {}
    for u in uns:
        by_slide.setdefault(int(u["slide"]), []).append(u)
    table, rewrites = [], []
    for r in rows:
        n, title, text = int(r["slide"]), r["title"], r["text"]
        body = text.splitlines()[1:] if text.splitlines() and text.splitlines()[0].lstrip("# ").strip() == title else text.splitlines()
        words = len(re.findall(r"\w+", " ".join(body)))
        verdicts, reasons = [], []
        figs = by_slide.get(n, [])
        if figs:
            verdicts += ["SHARPEN", "TBD"]
            reasons.append(f"{len(figs)} figure(s) without a source")
        if len(title.split()) <= 4 and not numbers_in(title) and not DECISION_RE.search(title):
            verdicts.append("SHARPEN")
            reasons.append("title is a label, not an assertion")
        if words > 60:
            verdicts.append("SHARPEN")
            reasons.append(f"body has {words} words (limit 60)")
        verdicts = list(dict.fromkeys(verdicts)) or ["KEEP"]
        tbd_items = "; ".join(f"{u['number']} → {DECK_TBD}" for u in figs) or "—"
        table.append(f"| {n} | {title.replace('|', '/')} | {' + '.join(f'[{v}]' for v in verdicts)} | "
                     f"{'; '.join(reasons) or 'passes the mechanical checks'} | {DECK_TBD if 'SHARPEN' in verdicts else '—'} | "
                     f"{'mark unsourced figures' if figs else '—'} | {tbd_items} | — |")
        if verdicts != ["KEEP"]:
            lines = []
            for ln in body:
                hits = [u["number"] for u in figs if u["context"] and ln.strip()[:160] == u["context"]]
                pos, out_ln = 0, ""
                for num in hits:                          # replace each figure once, left to right
                    i = ln.find(num, pos)
                    if i < 0:
                        continue
                    out_ln += ln[pos:i] + DECK_TBD
                    pos = i + len(num)
                lines.append(out_ln + ln[pos:])
            rewrites += [f"Slide {n} — {DECK_TBD if 'title is a label, not an assertion' in reasons else title}",
                         f"So what: {DECK_TBD}", "Body:"] + lines + \
                        [f"Source: {DECK_TBD if figs else 'as on the slide'}", ""]
    ask = any(DECISION_RE.search(r["text"]) for r in rows[:2])
    out = [f"# Deck Redline — {deck_name} — draft", "",
           "> Script draft — the model step was not run (draft only): mechanical checks from the rubric "
           "(sources, title form, density); judgement on storyline and so-what is an owned marker.", "",
           "## Bottom line", f"{DECK_TBD} — the decision the deck asks for, and whether the deck as found earns it.", "",
           "## Redline table",
           "| Slide # | Original title | Verdict | Reason | New title (assertion) | Body change (summary) | Source / `[TBD]` items | Moves to |",
           "|---|---|---|---|---|---|---|---|"] + table + ["", "## Rewritten slides", "", "```"] + \
          (rewrites or ["No slide needs a mechanical rewrite."]) + ["```", "", "## Unsourced numbers to baseline",
                                                                  "| Slide | Number | Context | Owner |", "|---|---|---|---|"]
    out += [f"| {u['slide']} | {u['number']} | {u['context'].replace('|', '/')} | {DECK_TBD} |" for u in uns] or \
        ["| — | none | every figure names a source | — |"]
    out += ["", "## Slides to add", "| Beat missing (storyline checklist) | Proposed slide | Why |", "|---|---|---|"]
    out.append("| Decision ask | none — slides 1–2 state it | — |" if ask else
               f"| Decision ask | Slide 1: the decision as an assertion, with owner and date | {DECK_TBD} |")
    return "\n".join(out) + "\n"


def apply(redline: Path, ledger: Path) -> None:
    if not redline.is_file():
        die(f"file not found: {redline}")
    if not ledger.is_file():
        die(f"ledger not found: {ledger}")
    verdicts = {}
    for ln in redline.read_text(encoding="utf-8", errors="replace").splitlines():
        if not ln.startswith("|"):
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if len(cells) < 3 or not re.fullmatch(r"\d+", cells[0]):
            continue
        vs = VERDICT_RE.findall(ln)
        if not vs:
            continue
        verdict = " + ".join(f"[{v}]" for v in dict.fromkeys(vs))
        reason = cells[3] if len(cells) > 3 else ""
        verdicts[int(cells[0])] = (verdict, reason)
    with ledger.open(encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    n = 0
    for r in rows:
        v = verdicts.get(int(r["slide"])) if r.get("slide", "").isdigit() else None
        if v:
            r["verdict"], r["reason"] = v
            n += 1
    with ledger.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LEDGER_COLS); w.writeheader(); w.writerows(rows)
    missing = [r["slide"] for r in rows if not r["verdict"]]
    print(f"applied {n} verdicts to {ledger}; {len(missing)} slides without verdict"
          + (f": {missing[:10]}" if missing else ""))


def main() -> int:
    ap = argparse.ArgumentParser(description="Deck -> slide ledger + unsourced numbers; --apply verdicts from redline.md")
    ap.add_argument("deck", nargs="?", help="deck.pptx, deck.md or deck.txt")
    ap.add_argument("--ledger", required=True, help="ledger.csv (written, or updated with --apply)")
    ap.add_argument("--unsourced", default="", help="unsourced-numbers.csv (build mode)")
    ap.add_argument("--apply", metavar="REDLINE_MD", help="fill verdict/reason in the ledger from redline.md")
    ap.add_argument("--redline-out", default="", help="build mode: draft redline path (written with --draft-only)")
    ap.add_argument("--draft-only", action="store_true", help="no model step: write --redline-out from the ledger")
    a = ap.parse_args()
    if a.apply:
        apply(Path(a.apply), Path(a.ledger))
        return 0
    if not a.deck or not a.unsourced:
        die("build mode needs a deck and --unsourced")
    build(Path(a.deck), Path(a.ledger), Path(a.unsourced),
          Path(a.redline_out) if a.draft_only and a.redline_out else None)
    return 0


if __name__ == "__main__":
    main()
